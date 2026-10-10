#!/usr/bin/env python3
"""Source Data: one workbook per figure, one sheet per panel, with the numbers plotted.

Each sheet opens with the panel it belongs to and where its numbers come from. Values the
figure derives (Wilson intervals, fold enrichments, hypergeometric P values, expected
counts) are written as the figure computes them, next to the inputs they come from.
"""

from __future__ import annotations

import json

from openpyxl import Workbook
from openpyxl.styles import Font
from scipy.stats import hypergeom

import figlib as fl
import nature_style as ns

OUT = ns.MANUSCRIPT / "source_data"
BOLD = Font(bold=True)
ITALIC = Font(italic=True, color="555555")


def sheet(wb, name, title, source, header, rows):
    ws = wb.create_sheet(name[:31])
    ws.append([title])
    ws["A1"].font = BOLD
    ws.append([f"Source: {source}"])
    ws["A2"].font = ITALIC
    ws.append([])
    ws.append(header)
    for c in ws[4]:
        c.font = BOLD
    for r in rows:
        ws.append([_num(v) for v in r])
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[3:]) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 60)
    return ws


def _num(v):
    if isinstance(v, str):
        try:
            f = float(v)
            return int(f) if f.is_integer() and "." not in v and "e" not in v.lower() else f
        except ValueError:
            return v
    return v


def book(name, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        sheet(wb, *s)
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"Source_Data_{name}.xlsx"
    wb.properties.creator = "TCMScience manuscript/code"
    wb.save(path)
    fl.deterministic_zip(path)
    return path


def wilson_cols(k, n):
    lo, hi = ns.wilson(int(k), int(n))
    return [round(int(k) / int(n), 6), round(lo, 6), round(hi, 6)]


def fig1():
    inv = fl.load_json("inventory.json")
    return book("Fig1", [
        ("a_inventory", "Fig. 1a: counts shown in the capability plane and kernel",
         "data/extracted/inventory.json (bioagent.providers, bioagent.tools, "
         "psh.compiler.diagnostics, BioScience-Harness/data/connector_live_verification.csv)",
         ["quantity", "value"],
         [["public sources", inv["public_sources"]],
          *[[f"public sources: {k}", v] for k, v in inv["public_sources_by_group"].items()],
          ["typed operations", inv["typed_operations"]],
          ["operations verified live", inv["operations_live_verified"]],
          ["verification dates", " to ".join(inv["verification_dates"])],
          ["native tools", inv["native_tools"]],
          ["native tool domains", inv["native_tool_domains"]],
          ["compiler diagnostics", inv["compiler_diagnostics"]],
          ["claim reason codes", inv["claim_reason_codes"]]]),
        ("c_tools_by_domain", "Fig. 1c: native tools by domain", "data/extracted/inventory.json",
         ["domain", "tools"], sorted(inv["native_tools_by_domain"].items(), key=lambda x: -x[1])),
        ("c_diagnostics", "Fig. 1c: compiler diagnostics by family",
         "data/extracted/inventory.json (psh.compiler.diagnostics.REGISTRY)",
         ["family", "codes"],
         sorted(inv["compiler_diagnostics_by_family"].items(), key=lambda x: -x[1])),
        ("claim_codes", "Claim reason codes (Figs. 1-2)",
         "bioagent.contracts.candidate_claim.CLAIM_REASONS", ["code", "meaning"],
         list(inv["claim_reasons"].items())),
    ])


def fig2():
    lic = fl.rows("licensing_matrix.csv")
    cert = fl.rows("max_certainty.csv")
    sup = fl.rows("compiler_supports.csv")
    claims = fl.rows("e2e_claims.csv")
    return book("Fig2", [
        ("a_licensing", "Fig. 2a: best grade each design can reach for each claim kind",
         "psh.sir.values.LICENSING via data/extracted/licensing_matrix.csv",
         ["claim_kind", "design", "grade"], [[r["claim_kind"], r["design"], r["grade"]] for r in lic]),
        ("a_certainty", "Fig. 2a: strongest wording a single study of each design licenses",
         "psh.sir.values.MAX_CERTAINTY", ["design", "max_certainty"],
         [[r["design"], r["max_certainty"]] for r in cert]),
        ("a_compile_time_table", "Compile-time table (EVIDENCE103); checked to be a subset of "
         "the direct cells of Fig. 2a", "psh.workflow.compiler._SUPPORTS",
         ["claim_type", "design", "predictive", "licensed"],
         [[r["claim_type"], r["design"], r["predictive"], r["licensed"]] for r in sup]),
        ("c_claims", "Fig. 2c: drafted claims and verdicts in the three end-to-end cases",
         "cases 1-2 re-run here (scripts/run_end_to_end_cases.py); case 3 as recorded in "
         "docs/end-to-end-cases.md:90-93", ["case", "order", "claim", "kind", "allowed",
                                            "codes", "origin"],
         [[r["case"], r["order"], r["text"], r["kind"], r["allowed"], r["codes"], r["origin"]]
          for r in claims]),
    ])


def fig3():
    conf = fl.rows("ablation_configurations.csv")
    km = fl.rows("ablation_kill_matrix.csv")
    un = fl.rows("ablation_unique.csv")
    rounds = fl.rows("ablation_rounds.csv")
    cases = fl.rows("ablation_cases.csv")
    base = fl.rows("ablation_base_cases.csv")
    src = "BioScience-Harness/benchmarks/ablation/results.json"
    return book("Fig3", [
        ("a_base_cases", "Fig. 3a: the ten base outputs", "bioagent.benchmarks.ablation_corpus",
         ["case", "language", "claim_kind", "source_design", "statement"],
         [[r["case"], r["language"], r["claim_kind"], r["source_design"], r["statement"]]
          for r in base]),
        ("b_configurations", "Fig. 3b: erroneous outputs released by configuration "
         "(Wilson 95% CI)", src,
         ["configuration", "gates", "errors_released", "mutants", "rate", "ci_low", "ci_high",
          "false_refusals", "base_cases"],
         [[r[k] for k in ("configuration", "gates", "errors_released", "mutants", "rate",
                          "ci_low", "ci_high", "false_refusals", "base_cases")] for r in conf]),
        ("c_kill_matrix", "Fig. 3c: share of each error class's mutants refused", src,
         list(km[0].keys()), [list(r.values()) for r in km]),
        ("d_leave_one_out", "Fig. 3d: errors the full stack refuses and releases once the gate "
         "is off", src, ["gate", "layer", "unique", "mutants"],
         [[r["gate"], r["layer"], r["unique"], r["mutants"]] for r in un]),
        ("e_rounds", "Fig. 3e: errors the full stack released in each round (Wilson 95% CI)",
         "docs/governance-ablation.md:94,108; results.json", ["stage", "errors_released",
                                                              "mutants", "rate", "ci_low",
                                                              "ci_high", "note"],
         [[r["stage"], r["errors_released"], r["mutants"],
           *wilson_cols(r["errors_released"], r["mutants"]), r["note"]] for r in rounds]),
        ("mutants", "Every case: which gates refuse it", src,
         ["case", "base", "error_class", "gold", "language", "refused_by",
          "refused_by_output_after_ingest"],
         [[r[k] for k in ("case", "base", "error_class", "gold", "language", "refused_by",
                          "refused_by_output_after_ingest")] for r in cases]),
    ])


def fig4():
    conf = fl.rows("evidence_typing_confusion_test.csv")
    des = fl.rows("evidence_typing_designs.csv")
    over = fl.rows("evidence_typing_overall.csv")
    fields = fl.rows("evidence_typing_fields_test.csv")
    manual = fl.rows("evidence_typing_manual_check.csv")
    samp = fl.rows("evidence_typing_sampling.csv")
    src = "BioScience-Harness/benchmarks/evidence_typing/results.json (with its baseline)"
    return book("Fig4", [
        ("a_sampling", "Fig. 4a: strata of the frozen sample", "benchmarks/evidence_typing/"
         "sampling.json", list(samp[0].keys()), [list(r.values()) for r in samp]),
        ("b_confusion_test", "Fig. 4b: reference (rows) against the design the rules read "
         "(columns), test split", src, list(conf[0].keys()), [list(r.values()) for r in conf]),
        ("c_designs", "Fig. 4c: precision and recall by design, before and after the rule "
         "change (Wilson 95% CI)", src, list(des[0].keys()), [list(r.values()) for r in des]),
        ("d_overall", "Fig. 4d: readings correct, studies typed and negatives typed",
         src, list(over[0].keys()), [list(r.values()) for r in over]),
        ("e_fields", "Fig. 4e: population, comparator and outcome filled (test split)", src,
         list(fields[0].keys()), [list(r.values()) for r in fields]),
        ("e_manual_check", "Fig. 4e: 30 filled values read by hand",
         "benchmarks/evidence_typing/manual_check.json", list(manual[0].keys()),
         [list(r.values()) for r in manual]),
    ])


def fig5():
    c = fl.counts()
    counts = fl.rows("np_counts.csv")
    prot = fl.rows("np_screening_proteins.csv")
    dis = fl.rows("np_disease_sets.csv")
    th = fl.rows("np_thresholds.csv")
    tp = fl.rows("np_threshold_pathways.csv")
    ca = []
    for name, prot_n in (("whole annotation", c["reactome_background"]),
                         ("assayed proteins", c["assayed_background"])):
        expected = c["ca_pathway_members"] * c["assayed_hits"] / prot_n
        ca.append([name, prot_n, c["assayed_hits"], c["ca_pathway_members"],
                   c["ca_pathway_hits"], round(expected, 4),
                   round(c["ca_pathway_hits"] / expected, 4)])
    drows = []
    for r in dis:
        M, n = (12155, 237) if "whole" in r["background"] else (391, 237)
        N, k = int(r["disease_genes"]), int(r["overlap"])
        drows.append([r["label"], r["background"], N, k, r["fold"], round(k / (n * N / M), 4),
                      r["p_hypergeometric"], float(f"{hypergeom.sf(k - 1, M, n, N):.4g}"),
                      r["p_permutation"], r["source"]])
    prows = [[r["family"], r["protein"], r["active"], r["tested"],
              *wilson_cols(r["active"], r["tested"]), r["assay_origin"], r["source"]]
             for r in prot]
    return book("Fig5", [
        ("a_counts", "Fig. 5a and b: the case study's counts, each with the line it is quoted "
         "from", "data/curated/np_counts.csv", ["key", "value", "unit", "description", "source",
                                                "quote"],
         [[r[k] for k in ("key", "value", "unit", "description", "source", "quote")]
          for r in counts]),
        ("b_ca_pathway", "Fig. 5b: the carbonic-anhydrase pathway under each background "
         "(expected hits and fold computed from the counts)", "data/curated/np_counts.csv",
         ["background", "proteins", "hits_in_background", "pathway_members", "pathway_hits",
          "expected_hits", "fold"], ca),
        ("c_disease_genes", "Fig. 5c: measured targets against type 2 diabetes genes; fold "
         "and one-sided hypergeometric P recomputed from the counts",
         "data/curated/np_disease_sets.csv",
         ["definition", "background", "disease_genes", "overlap", "fold_reported",
          "fold_recomputed", "p_reported", "p_recomputed", "p_permutation", "source"], drows),
        ("d_screening_proteins", "Fig. 5d: active compound-protein tests per protein (Wilson "
         "95% CI); all 472 proteins: 3,357/51,304 active", "data/curated/np_screening_proteins.csv",
         ["family", "protein", "active", "tested", "share", "ci_low", "ci_high", "assay_origin",
          "source"], prows),
        ("e_thresholds", "Fig. 5e: the screening analysis at each minimum coverage",
         "data/curated/np_thresholds.csv", ["min_constituents", "proteins", "pathways_tested",
                                            "active_percent", "pathways_passing", "source"],
         [[r[k] for k in ("min_constituents", "proteins", "pathways_tested", "active_percent",
                          "pathways_passing", "source")] for r in th]),
        ("e_pathways", "Fig. 5e: pathways passing BH and 10,000 permutations at each threshold "
         "(q where the source states it)", "data/curated/np_threshold_pathways.csv",
         ["pathway", "carrier", "min_constituents", "q", "source"],
         [[r[k] for k in ("pathway", "carrier", "min_constituents", "q", "source")] for r in tp]),
    ])


def fig6():
    hold = fl.rows("robustness_holdouts.csv")
    pk = fl.rows("berberine_pk.csv")
    expo = fl.rows("berberine_exposure.csv")
    summ = fl.load_json("berberine_summary.json")
    routes = fl.rows("routes.csv")
    return book("Fig6", [
        ("a_holdouts", "Fig. 6a: the released hypotheses with evidence held out",
         "data/curated/robustness_holdouts.csv (docs/falsifiable-studies.md:58-60)",
         ["pathway", "holdout", "q", "status", "source"],
         [[r[k] for k in ("pathway", "holdout", "q", "status", "source")] for r in hold]),
        ("b_pk_ratios", "Fig. 6b: ratios of berberine exposure (90% CI, delta method on the "
         "log scale, Welch-Satterthwaite df)", "docs/studies/gqd_berberine.json",
         list(pk[0].keys()), [list(r.values()) for r in pk]),
        ("c_groups", "Fig. 6c: plasma Cmax and verdict counts per group",
         "docs/studies/gqd_berberine.json",
         ["group", "cmax_ng_ml", "cmax_nM_total", "plausible", "depends_on_unknowns",
          "implausible"],
         [[g, v["cmax_ng_ml"], v["cmax_nM_total"], v["verdicts"].get("plausible", 0),
           v["verdicts"].get("depends_on_unknowns", 0), v["verdicts"].get("implausible", 0)]
          for g, v in summ["groups"].items()]),
        ("c_assays", "Fig. 6c: ToxCast actives not ruled out, with the exposure ratio range "
         "R = C x fu / AC50 for fu from 0.01 to 1", "docs/studies/gqd_berberine.json",
         list(expo[0].keys()), [list(r.values()) for r in expo]),
        ("d_routes", "Fig. 6d: the eight routes and their status",
         "data/curated/routes.csv (docs/falsifiable-studies.md:116-123)",
         ["route", "name", "status", "shown", "source"],
         [[r[k] for k in ("route", "name", "status", "shown", "source")] for r in routes]),
    ])


def fig7():
    pred = fl.rows("inquiry_predictions.csv")
    steps = fl.rows("inquiry_steps.csv")
    opts = fl.rows("inquiry_options.csv")
    summ = fl.rows("inquiry_summary.csv")
    design = fl.load_json("inquiry_design.json")
    src = ("BioScience-Harness/scripts/run_inquiry.py --planted all, replayed through "
           "psh.scientist.inquiry (code/extract_data.py)")
    return book("Fig7", [
        ("b_predictions", "Fig. 7b: sealed P(outcome | explanation)", src,
         ["analysis", "explanation", "outcome", "p"],
         [[r["analysis"], r["explanation"], r["outcome"], r["p"]] for r in pred]),
        ("b_design", "Fig. 7b: analyses, costs (network-pharmacology runs) and priors", src,
         ["kind", "id", "cost_or_prior", "outcomes_or_role"],
         [["analysis", a["id"], a["cost"], ";".join(a["outcomes"])] for a in design["analyses"]]
         + [["explanation", e["id"], e["prior"], e["role"]] for e in design["explanations"]]),
        ("c_posteriors", "Fig. 7c: belief after each step in each planted world", src,
         list(steps[0].keys()), [list(r.values()) for r in steps]),
        ("d_options", "Fig. 7d: expected bits (and bits per run) of every analysis at each "
         "decision; chosen = 1 marks the one the engine ran", src,
         list(opts[0].keys()), [list(r.values()) for r in opts]),
        ("summary", "Verdicts, runs used and claims released", src,
         list(summ[0].keys()), [list(r.values()) for r in summ]),
    ])


def main():
    return [fig1(), fig2(), fig3(), fig4(), fig5(), fig6(), fig7()]


if __name__ == "__main__":
    for p in main():
        print(p)
